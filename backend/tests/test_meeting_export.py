import io
import json
import re
import tempfile
import zlib
from datetime import date
from urllib.parse import unquote

import pytest
from docx import Document
from sqlalchemy import select

from app.db.database import SessionLocal
from app.models.action_item import ActionItem
from app.models.decision import Decision
from app.models.meeting import Meeting
from app.models.transcript_segment import TranscriptSegment
from app.services import export_service
from tests.test_transcription import auth_header, register_and_login

UNICODE_TITLE = 'Café "Q4" planning / नमस्ते'

MIME_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain; charset=utf-8",
    "json": "application/json; charset=utf-8",
}


def new_meeting(client, token, title="Launch sync", description=None):
    payload = {"title": title, "meeting_date": "2026-10-09T10:00:00Z"}
    if description is not None:
        payload["description"] = description
    created = client.post("/meetings", headers=auth_header(token), json=payload)
    assert created.status_code == 201, created.text
    return created.json()["id"]


def add_segment(client, token, meeting_id, speaker, start, text):
    created = client.post(
        f"/meetings/{meeting_id}/transcript",
        headers=auth_header(token),
        json={"speaker": speaker, "start_time": start, "end_time": start + 4, "text": text},
    )
    assert created.status_code == 201, created.text


def populated_meeting(client, token, title="Launch sync"):
    meeting_id = new_meeting(client, token, title, description="Weekly launch check-in.")
    # Added out of order; exports must follow the timestamps.
    add_segment(client, token, meeting_id, "Speaker 2", 20.0, "Second: the beta is ready.")
    add_segment(client, token, meeting_id, "Speaker 1", 3.0, "First: welcome everyone.")
    add_segment(client, token, meeting_id, "Speaker 1", 75.5, 'Third: ship it with "quotes" & <tags>.')

    saved = client.put(
        f"/meetings/{meeting_id}/summary",
        headers=auth_header(token),
        json={"summary": "The team agreed to ship the beta."},
    )
    assert saved.status_code == 200, saved.text
    item = client.post(
        f"/meetings/{meeting_id}/action-items",
        headers=auth_header(token),
        json={"task": "Send the release notes", "assignee": "Priya", "due_date": "2026-10-14"},
    )
    assert item.status_code == 201, item.text
    done = client.post(
        f"/meetings/{meeting_id}/action-items",
        headers=auth_header(token),
        json={"task": "Book the retro", "status": "completed"},
    )
    assert done.status_code == 201, done.text
    decision = client.post(
        f"/meetings/{meeting_id}/decisions",
        headers=auth_header(token),
        json={"decision": "Launch the beta on Tuesday", "timestamp": 42.0, "context": "Review passed."},
    )
    assert decision.status_code == 201, decision.text
    return meeting_id


def export(client, token, meeting_id, export_format, **params):
    return client.get(
        f"/meetings/{meeting_id}/export",
        headers=auth_header(token),
        params={"format": export_format, **{key: str(value).lower() for key, value in params.items()}},
    )


def docx_text(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def pdf_page_count(content: bytes) -> int:
    return len(re.findall(rb"/Type\s*/Page\b", content))


def pdf_stream_text(content: bytes) -> bytes:
    """Content streams of a PDF drawn with a core font, where text is literal."""
    text = b""
    for match in re.finditer(rb"stream\r?\n(.*?)\r?\nendstream", content, re.S):
        try:
            text += zlib.decompress(match.group(1))
        except zlib.error:
            text += match.group(1)
    return text


def snapshot(meeting_id):
    db = SessionLocal()
    try:
        meeting = db.get(Meeting, meeting_id)
        rows = {
            "meeting": {column.name: getattr(meeting, column.name) for column in Meeting.__table__.columns},
        }
        for model in (ActionItem, Decision, TranscriptSegment):
            rows[model.__tablename__] = [
                {column.name: getattr(row, column.name) for column in model.__table__.columns}
                for row in db.scalars(select(model).where(model.meeting_id == meeting_id).order_by(model.id))
            ]
        return rows
    finally:
        db.close()


# Formats, MIME types and filenames


@pytest.mark.parametrize("export_format", ["pdf", "docx", "txt", "json"])
def test_each_format_downloads_with_mime_type_and_safe_filename(client, export_format):
    _, token = register_and_login(client, "Owner", f"export-formats-{export_format}@example.com")
    meeting_id = populated_meeting(client, token, UNICODE_TITLE)

    response = export(client, token, meeting_id, export_format)

    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == MIME_TYPES[export_format]
    assert response.headers["cache-control"] == "no-store"
    disposition = response.headers["content-disposition"]
    assert disposition.startswith("attachment; ")
    ascii_name = re.search(r'filename="([^"]+)"', disposition).group(1)
    assert ascii_name == f"Cafe-Q4-planning-2026-10-09.{export_format}"
    unicode_name = unquote(re.search(r"filename\*=UTF-8''(\S+)", disposition).group(1))
    assert unicode_name == f"Café Q4 planning नमस्ते-2026-10-09.{export_format}"
    assert not re.search(r'[/\\:*?"<>|]', unicode_name)
    assert response.content


def test_default_format_is_pdf(client):
    _, token = register_and_login(client, "Owner", "export-default@example.com")
    meeting_id = new_meeting(client, token)

    response = client.get(f"/meetings/{meeting_id}/export", headers=auth_header(token))

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"


def test_filename_falls_back_when_title_has_no_usable_characters():
    ascii_name, unicode_name = export_service.export_filenames("///", 12, None, "txt")
    assert (ascii_name, unicode_name) == ("meeting-12.txt", "meeting-12.txt")
    ascii_name, unicode_name = export_service.export_filenames("会议", 12, None, "txt")
    assert ascii_name == "meeting-12.txt"
    assert unicode_name == "会议.txt"
    ascii_name, _ = export_service.export_filenames("CON", 3, None, "pdf")
    assert ascii_name == "meeting-3.pdf"
    ascii_name, _ = export_service.export_filenames("x" * 300, 3, None, "pdf")
    assert len(ascii_name) <= export_service.MAX_FILENAME_STEM + len(".pdf")


def test_unknown_format_is_rejected(client):
    _, token = register_and_login(client, "Owner", "export-bad-format@example.com")
    meeting_id = new_meeting(client, token)

    assert export(client, token, meeting_id, "html").status_code == 422


def test_export_needs_at_least_one_section(client):
    _, token = register_and_login(client, "Owner", "export-no-sections@example.com")
    meeting_id = new_meeting(client, token)

    response = export(
        client, token, meeting_id, "txt",
        details=False, summary=False, action_items=False, decisions=False, transcript=False,
    )

    assert response.status_code == 422
    assert response.json()["detail"] == "Select at least one section to export."


# Content


def test_json_export_is_valid_utf8_with_ordered_records(client):
    _, token = register_and_login(client, "Owner", "export-json@example.com")
    meeting_id = populated_meeting(client, token, UNICODE_TITLE)

    response = export(client, token, meeting_id, "json")
    data = json.loads(response.content.decode("utf-8"))

    assert data["format_version"] == 1
    assert data["sections"] == ["details", "summary", "action_items", "decisions", "transcript"]
    assert data["meeting"]["id"] == meeting_id
    assert data["meeting"]["title"] == UNICODE_TITLE
    assert data["meeting"]["meeting_date"] == "2026-10-09T10:00:00Z"
    assert data["meeting"]["has_recording"] is False
    assert "recording_path" not in data["meeting"]
    assert "user_id" not in data["meeting"]

    assert data["summary"] == {
        "text": "The team agreed to ship the beta.",
        "source": "manual",
        "method": None,
        "key_points": [],
    }
    assert [segment["text"] for segment in data["transcript"]] == [
        "First: welcome everyone.",
        "Second: the beta is ready.",
        'Third: ship it with "quotes" & <tags>.',
    ]
    assert [segment["speaker"] for segment in data["transcript"]] == ["Speaker 1", "Speaker 2", "Speaker 1"]
    assert data["transcript"][0]["start_time"] == 3.0
    assert all(isinstance(segment["id"], int) for segment in data["transcript"])

    pending, completed = data["action_items"]
    assert pending["task"] == "Send the release notes"
    assert pending["assignee"] == "Priya"
    assert pending["due_date"] == "2026-10-14"
    assert pending["status"] == "pending"
    assert pending["completed"] is False
    assert completed["status"] == "completed"
    assert completed["completed"] is True
    assert completed["assignee"] is None
    assert completed["due_date"] is None

    assert data["decisions"] == [
        {
            "id": data["decisions"][0]["id"],
            "decision": "Launch the beta on Tuesday",
            "timestamp": 42.0,
            "context": "Review passed.",
            "source": "manual",
            "created_at": data["decisions"][0]["created_at"],
        }
    ]
    assert "नमस्ते".encode() in response.content


def test_txt_export_lists_every_section(client):
    _, token = register_and_login(client, "Owner", "export-txt@example.com")
    meeting_id = populated_meeting(client, token, UNICODE_TITLE)

    text = export(client, token, meeting_id, "txt").content.decode("utf-8")

    assert text.startswith(UNICODE_TITLE + "\n")
    for expected in (
        "MEETING DETAILS",
        "Date: October 9, 2026 at 10:00 UTC",
        "Recording: Not available",
        "Weekly launch check-in.",
        "SUMMARY",
        "The team agreed to ship the beta.",
        "(Written or edited by a user.)",
        "ACTION ITEMS",
        "1. [Pending] Send the release notes",
        "Assignee: Priya | Due: 2026-10-14",
        "2. [Completed] Book the retro",
        "DECISIONS",
        "1. [00:42] Launch the beta on Tuesday",
        "Context: Review passed.",
        "TRANSCRIPT",
        "[00:03 - 00:07] Speaker 1: First: welcome everyone.",
        '[01:15 - 01:19] Speaker 1: Third: ship it with "quotes" & <tags>.',
    ):
        assert expected in text
    assert text.index("First: welcome") < text.index("Second: the beta") < text.index("Third: ship")


def test_docx_export_opens_as_a_word_document(client):
    _, token = register_and_login(client, "Owner", "export-docx@example.com")
    meeting_id = populated_meeting(client, token, UNICODE_TITLE)

    content = export(client, token, meeting_id, "docx").content
    document = Document(io.BytesIO(content))
    text = docx_text(content)

    assert document.core_properties.title == UNICODE_TITLE
    headings = [p.text for p in document.paragraphs if p.style.name.startswith(("Heading", "Title"))]
    assert headings == [UNICODE_TITLE, "Meeting details", "Description", "Summary", "Action items", "Decisions", "Transcript"]
    assert "Task | Assignee | Due date | Status | Time" in text
    assert "Send the release notes | Priya | 2026-10-14 | Pending | \u2014" in text
    assert "00:42 | Launch the beta on Tuesday | Review passed." in text
    assert "[00:03 - 00:07] Speaker 1: First: welcome everyone." in text
    assert text.index("First: welcome") < text.index("Second: the beta") < text.index("Third: ship")


def test_pdf_export_is_a_valid_document(client):
    _, token = register_and_login(client, "Owner", "export-pdf@example.com")
    meeting_id = populated_meeting(client, token, UNICODE_TITLE)

    content = export(client, token, meeting_id, "pdf").content

    assert content.startswith(b"%PDF-")
    assert content.rstrip().endswith(b"%%EOF")
    assert pdf_page_count(content) >= 1


def test_pdf_text_is_drawn_in_order_without_a_unicode_font(client, monkeypatch):
    monkeypatch.setattr(export_service, "_available_faces", lambda: [])
    _, token = register_and_login(client, "Owner", "export-pdf-core@example.com")
    meeting_id = populated_meeting(client, token, UNICODE_TITLE)

    content = export(client, token, meeting_id, "pdf").content
    streams = pdf_stream_text(content)

    assert content.startswith(b"%PDF-")
    for expected in (b"Meeting details", b"Send the release notes", b"Launch the beta on Tuesday", b"Speaker 2"):
        assert expected in streams
    assert streams.index(b"First: welcome") < streams.index(b"Second: the beta") < streams.index(b"Third: ship")


def test_long_transcript_spans_many_pdf_pages(client):
    _, token = register_and_login(client, "Owner", "export-long@example.com")
    meeting_id = new_meeting(client, token)
    db = SessionLocal()
    try:
        db.add_all(
            TranscriptSegment(
                meeting_id=meeting_id,
                speaker=f"Speaker {index % 3 + 1}",
                start_time=index * 6.0,
                end_time=index * 6.0 + 5,
                text=f"Segment {index}. " + "This sentence keeps the transcript long enough to wrap. " * 4,
                source="ai",
            )
            for index in range(400)
        )
        db.commit()
    finally:
        db.close()

    short = export(client, token, meeting_id, "pdf", transcript=False).content
    long = export(client, token, meeting_id, "pdf").content
    text = export(client, token, meeting_id, "txt").content.decode("utf-8")

    assert pdf_page_count(short) == 1
    assert pdf_page_count(long) > 20
    assert text.count("Speaker ") >= 400
    assert "[39:54 - 39:59] Speaker 1: Segment 399." in text


def test_very_long_action_item_still_renders(client):
    _, token = register_and_login(client, "Owner", "export-long-item@example.com")
    meeting_id = new_meeting(client, token)
    task = "Prepare the migration plan " * 120
    created = client.post(f"/meetings/{meeting_id}/action-items", headers=auth_header(token), json={"task": task})
    assert created.status_code == 201

    for export_format in ("pdf", "docx", "txt"):
        response = export(client, token, meeting_id, export_format, transcript=False)
        assert response.status_code == 200, export_format
    assert task.strip() in docx_text(export(client, token, meeting_id, "docx").content)


def test_unchecked_sections_are_left_out(client):
    _, token = register_and_login(client, "Owner", "export-sections@example.com")
    meeting_id = populated_meeting(client, token)

    data = export(client, token, meeting_id, "json", transcript=False, decisions=False, details=False).json()
    text = export(client, token, meeting_id, "txt", summary=False, action_items=False).content.decode()
    docx = docx_text(export(client, token, meeting_id, "docx", transcript=False, summary=False).content)

    assert data["sections"] == ["summary", "action_items"]
    assert set(data) == {"format_version", "generated_at", "sections", "meeting", "summary", "action_items"}
    assert data["meeting"] == {"id": meeting_id, "title": "Launch sync"}

    assert "SUMMARY" not in text and "ACTION ITEMS" not in text
    assert "MEETING DETAILS" in text and "DECISIONS" in text and "TRANSCRIPT" in text

    assert "First: welcome" not in docx and "The team agreed" not in docx
    assert "Send the release notes" in docx


def test_empty_meeting_exports_cleanly(client):
    _, token = register_and_login(client, "Owner", "export-empty@example.com")
    meeting_id = new_meeting(client, token, "Empty")

    data = export(client, token, meeting_id, "json").json()
    text = export(client, token, meeting_id, "txt").content.decode()
    docx = docx_text(export(client, token, meeting_id, "docx").content)
    pdf = export(client, token, meeting_id, "pdf").content

    assert data["summary"] == {"text": None, "source": None, "method": None, "key_points": []}
    assert data["action_items"] == [] and data["decisions"] == [] and data["transcript"] == []
    assert data["meeting"]["description"] is None
    for expected in ("No summary has been saved for this meeting.", "No action items.", "No decisions.", "No transcript."):
        assert expected in text
        assert expected in docx
    assert "None" not in text and "null" not in text
    assert pdf.startswith(b"%PDF-")


def test_generated_insights_are_labelled_and_key_points_exported(client):
    _, token = register_and_login(client, "Owner", "export-insights@example.com")
    meeting_id = new_meeting(client, token)
    db = SessionLocal()
    try:
        meeting = db.get(Meeting, meeting_id)
        meeting.summary = "Generated summary."
        meeting.summary_source = "ai"
        meeting.key_points = [{"text": "Budget approved", "timestamp": 65.0}, "Legacy point"]
        db.add(ActionItem(meeting_id=meeting_id, task="Draft plan", timestamp=12.5, source="ai", due_date=date(2026, 10, 12)))
        db.commit()
    finally:
        db.close()

    data = export(client, token, meeting_id, "json").json()
    text = export(client, token, meeting_id, "txt").content.decode()

    assert data["summary"]["source"] == "ai"
    assert data["summary"]["method"] == "rule-based"
    assert data["summary"]["key_points"] == [
        {"text": "Budget approved", "timestamp": 65.0},
        {"text": "Legacy point", "timestamp": None},
    ]
    assert data["action_items"][0]["timestamp"] == 12.5
    assert data["action_items"][0]["source"] == "ai"
    assert "(Generated from the transcript by rule-based extraction.)" in text
    assert "  - Budget approved [01:05]" in text
    assert "Due: 2026-10-12 | Time: 00:12" in text


def test_special_and_control_characters_survive_every_format(client):
    _, token = register_and_login(client, "Owner", "export-special@example.com")
    meeting_id = new_meeting(client, token, "Ünïcödé — 会议")
    tricky = 'Tab\there, "quotes", <b>&amp;</b>, emoji 🎉, नमस्ते, 会议, bell\x07 end'
    add_segment(client, token, meeting_id, "Speaker 1", 1.0, tricky)

    data = json.loads(export(client, token, meeting_id, "json").content)
    text = export(client, token, meeting_id, "txt").content.decode("utf-8")
    docx = docx_text(export(client, token, meeting_id, "docx").content)
    pdf = export(client, token, meeting_id, "pdf")

    assert data["transcript"][0]["text"] == tricky
    assert '"quotes", <b>&amp;</b>, emoji 🎉, नमस्ते, 会议, bell end' in text
    assert '"quotes", <b>&amp;</b>, emoji 🎉, नमस्ते, 会议, bell end' in docx
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF-")


# Access, safety and side effects


def test_export_requires_authentication(client):
    _, token = register_and_login(client, "Owner", "export-auth@example.com")
    meeting_id = new_meeting(client, token)

    response = client.get(f"/meetings/{meeting_id}/export", params={"format": "json"})

    assert response.status_code == 401


def test_other_users_cannot_export_a_meeting(client):
    _, owner = register_and_login(client, "Owner", "export-owner@example.com")
    _, other = register_and_login(client, "Other", "export-other@example.com")
    meeting_id = populated_meeting(client, owner)

    for export_format in ("pdf", "json"):
        response = export(client, other, meeting_id, export_format)
        assert response.status_code == 404
        assert response.json() == {"detail": "Meeting not found"}
    assert export(client, owner, 999_999, "json").status_code == 404


def test_export_does_not_leak_paths_or_credentials(client, recordings_dir):
    user_id, token = register_and_login(client, "Owner", "export-secrets@example.com")
    meeting_id = populated_meeting(client, token)
    (recordings_dir / "secret-clip.webm").write_bytes(b"audio")
    db = SessionLocal()
    try:
        db.get(Meeting, meeting_id).recording_path = "recordings/secret-clip.webm"
        db.commit()
    finally:
        db.close()

    for export_format in ("json", "txt", "docx"):
        content = export(client, token, meeting_id, export_format).content
        if export_format == "docx":
            content = docx_text(content).encode()
        assert b"secret-clip" not in content
        assert str(recordings_dir).encode() not in content
        assert token.encode() not in content
        assert b"password123" not in content
        assert b"export-secrets@example.com" not in content
    assert export(client, token, meeting_id, "json").json()["meeting"]["has_recording"] is True


def test_missing_recording_file_does_not_block_export(client):
    _, token = register_and_login(client, "Owner", "export-missing-recording@example.com")
    meeting_id = populated_meeting(client, token)
    db = SessionLocal()
    try:
        db.get(Meeting, meeting_id).recording_path = "recordings/deleted.webm"
        db.commit()
    finally:
        db.close()

    for export_format in ("pdf", "docx", "txt", "json"):
        assert export(client, token, meeting_id, export_format).status_code == 200
    assert export(client, token, meeting_id, "json").json()["meeting"]["has_recording"] is False


def test_export_never_changes_stored_data(client, capture_transcription_jobs):
    _, token = register_and_login(client, "Owner", "export-readonly@example.com")
    meeting_id = populated_meeting(client, token)
    before = snapshot(meeting_id)

    for export_format in ("pdf", "docx", "txt", "json"):
        assert export(client, token, meeting_id, export_format).status_code == 200

    assert snapshot(meeting_id) == before
    assert capture_transcription_jobs == []
    intelligence = client.get(f"/meetings/{meeting_id}/intelligence", headers=auth_header(token)).json()
    assert intelligence["status"] is None


def test_export_is_built_in_memory(client, monkeypatch, recordings_dir):
    _, token = register_and_login(client, "Owner", "export-memory@example.com")
    meeting_id = populated_meeting(client, token)

    def no_temp_files(*_args, **_kwargs):
        raise AssertionError("export must not create temporary files")

    for name in ("mkstemp", "mkdtemp", "NamedTemporaryFile", "TemporaryFile", "SpooledTemporaryFile", "TemporaryDirectory"):
        monkeypatch.setattr(tempfile, name, no_temp_files)

    for export_format in ("pdf", "docx", "txt", "json"):
        assert export(client, token, meeting_id, export_format).status_code == 200
    assert list(recordings_dir.iterdir()) == []


def test_cors_exposes_the_download_filename(client):
    _, token = register_and_login(client, "Owner", "export-cors@example.com")
    meeting_id = new_meeting(client, token)

    response = client.get(
        f"/meetings/{meeting_id}/export",
        headers={**auth_header(token), "Origin": "http://localhost:5173"},
        params={"format": "txt"},
    )

    assert "content-disposition" in response.headers.get("access-control-expose-headers", "").lower()
