"""Downloadable meeting reports built from stored records.

Exports are read-only: they never regenerate insights, run transcription or
write to the database, and every file is rendered in memory.
"""

from __future__ import annotations

import io
import json
import logging
import os
import re
import unicodedata
from dataclasses import dataclass
from datetime import UTC, date, datetime
from enum import Enum
from functools import lru_cache
from pathlib import Path
from urllib.parse import quote

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.action_item import ActionItem
from app.models.decision import Decision
from app.models.meeting import Meeting
from app.models.transcript_segment import TranscriptSegment
from app.services.intelligence_service import METHOD, intelligence_snapshot
from app.services.recording_service import resolve_stored_recording

logger = logging.getLogger(__name__)

FORMAT_VERSION = 1
SECTIONS = ("details", "summary", "action_items", "decisions", "transcript")


class ExportFormat(str, Enum):
    PDF = "pdf"
    DOCX = "docx"
    TXT = "txt"
    JSON = "json"


MEDIA_TYPES = {
    ExportFormat.PDF: "application/pdf",
    ExportFormat.DOCX: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ExportFormat.TXT: "text/plain; charset=utf-8",
    ExportFormat.JSON: "application/json; charset=utf-8",
}


@dataclass(frozen=True)
class ExportOptions:
    details: bool = True
    summary: bool = True
    action_items: bool = True
    decisions: bool = True
    transcript: bool = True

    def included(self) -> list[str]:
        return [name for name in SECTIONS if getattr(self, name)]


@dataclass(frozen=True)
class ExportFile:
    content: bytes
    media_type: str
    filename: str
    content_disposition: str


def export_meeting(
    db: Session,
    meeting: Meeting,
    export_format: ExportFormat,
    options: ExportOptions,
    *,
    generated_at: datetime | None = None,
) -> ExportFile:
    data = build_export_data(db, meeting, options, generated_at=generated_at)
    content = _RENDERERS[export_format](data)
    ascii_name, unicode_name = export_filenames(meeting.title, meeting.id, _utc(meeting.meeting_date), export_format.value)
    return ExportFile(
        content=content,
        media_type=MEDIA_TYPES[export_format],
        filename=unicode_name,
        content_disposition=content_disposition(ascii_name, unicode_name),
    )


# --------------------------------------------------------------------------
# Data snapshot
# --------------------------------------------------------------------------


def build_export_data(
    db: Session,
    meeting: Meeting,
    options: ExportOptions,
    *,
    generated_at: datetime | None = None,
) -> dict:
    """Collects the selected sections as plain values.

    The stored recording path is deliberately left out; only whether a
    recording exists is reported.
    """
    data: dict = {
        "format_version": FORMAT_VERSION,
        "generated_at": _utc(generated_at or datetime.now(UTC)),
        "sections": options.included(),
        "meeting": {"id": meeting.id, "title": meeting.title},
    }

    if options.details:
        data["meeting"].update(
            {
                "description": meeting.description or None,
                "meeting_date": _utc(meeting.meeting_date),
                "duration_seconds": meeting.duration_seconds,
                "platform": meeting.platform,
                "status": meeting.status,
                "meeting_code": meeting.meeting_code,
                "has_recording": resolve_stored_recording(meeting.recording_path) is not None,
                "transcription_status": meeting.transcription_status,
                "transcription_language": meeting.transcription_language,
                "created_at": _utc(meeting.created_at),
                "updated_at": _utc(meeting.updated_at),
            }
        )

    if options.summary:
        snapshot = intelligence_snapshot(meeting)
        summary = (snapshot["summary"] or "").strip()
        data["summary"] = {
            "text": summary or None,
            "source": snapshot["summary_source"] if summary else None,
            "method": METHOD if summary and snapshot["summary_source"] == "ai" else None,
            "key_points": snapshot["key_points"],
        }

    if options.action_items:
        items = db.scalars(
            select(ActionItem)
            .where(ActionItem.meeting_id == meeting.id)
            .order_by(ActionItem.created_at.asc(), ActionItem.id.asc())
        ).all()
        data["action_items"] = [
            {
                "id": item.id,
                "task": item.task,
                "assignee": item.assignee or None,
                "due_date": item.due_date,
                "status": item.status,
                "completed": item.status == "completed",
                "timestamp": item.timestamp,
                "source": item.source,
                "created_at": _utc(item.created_at),
                "updated_at": _utc(item.updated_at),
            }
            for item in items
        ]

    if options.decisions:
        decisions = db.scalars(
            select(Decision)
            .where(Decision.meeting_id == meeting.id)
            .order_by(Decision.timestamp.asc(), Decision.id.asc())
        ).all()
        data["decisions"] = [
            {
                "id": decision.id,
                "decision": decision.decision,
                "timestamp": decision.timestamp,
                "context": decision.context or None,
                "source": decision.source,
                "created_at": _utc(decision.created_at),
            }
            for decision in decisions
        ]

    if options.transcript:
        segments = db.scalars(
            select(TranscriptSegment)
            .where(TranscriptSegment.meeting_id == meeting.id)
            .order_by(TranscriptSegment.start_time.asc(), TranscriptSegment.id.asc())
        ).all()
        data["transcript"] = [
            {
                "id": segment.id,
                "start_time": segment.start_time,
                "end_time": segment.end_time,
                "speaker": segment.speaker,
                "text": segment.text,
                "source": segment.source,
            }
            for segment in segments
        ]

    return data


# --------------------------------------------------------------------------
# Filenames
# --------------------------------------------------------------------------

_UNSAFE_FILENAME = re.compile(r'[<>:"/\\|?*\x00-\x1f\x7f]+')
_RESERVED_NAMES = {"con", "prn", "aux", "nul", *(f"com{i}" for i in range(1, 10)), *(f"lpt{i}" for i in range(1, 10))}
MAX_FILENAME_STEM = 80


def export_filenames(title: str | None, meeting_id: int, meeting_date: datetime | None, extension: str) -> tuple[str, str]:
    """Returns an ASCII fallback name and a Unicode name for the download."""
    suffix = f"-{meeting_date:%Y-%m-%d}" if meeting_date else ""

    readable = _UNSAFE_FILENAME.sub(" ", unicodedata.normalize("NFC", title or ""))
    readable = re.sub(r"\s+", " ", readable).strip(" .")[:MAX_FILENAME_STEM].strip(" .")

    ascii_stem = unicodedata.normalize("NFKD", readable).encode("ascii", "ignore").decode("ascii")
    ascii_stem = re.sub(r"[^A-Za-z0-9._-]+", "-", ascii_stem).strip("-._")
    ascii_stem = re.sub(r"-{2,}", "-", ascii_stem)

    fallback = f"meeting-{meeting_id}"
    if not ascii_stem or ascii_stem.lower() in _RESERVED_NAMES:
        ascii_stem = fallback
    if not readable or readable.lower() in _RESERVED_NAMES:
        readable = ascii_stem

    return f"{ascii_stem}{suffix}.{extension}", f"{readable}{suffix}.{extension}"


def content_disposition(ascii_name: str, unicode_name: str) -> str:
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(unicode_name, safe='')}"


# --------------------------------------------------------------------------
# Shared formatting
# --------------------------------------------------------------------------

_XML_INVALID = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\ufffe\uffff\ud800-\udfff]")
DASH = "\u2014"


def _utc(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    if value.tzinfo is None:
        return value.replace(tzinfo=UTC)
    return value.astimezone(UTC)


def _clean(text: str | None) -> str:
    if not text:
        return ""
    return _XML_INVALID.sub("", text.replace("\r\n", "\n").replace("\r", "\n").replace("\t", "    "))


def _clock(seconds: float | None) -> str:
    if seconds is None:
        return ""
    total = max(0, int(seconds))
    hours, rest = divmod(total, 3600)
    minutes, secs = divmod(rest, 60)
    return f"{hours}:{minutes:02d}:{secs:02d}" if hours else f"{minutes:02d}:{secs:02d}"


def _datetime_label(value: datetime | None) -> str:
    if value is None:
        return ""
    return f"{value:%B} {value.day}, {value:%Y} at {value:%H:%M} UTC"


def _date_label(value: date | None) -> str:
    return value.isoformat() if value else ""


def _duration_label(seconds: int | None) -> str:
    if not seconds:
        return ""
    hours, rest = divmod(int(seconds), 3600)
    minutes, secs = divmod(rest, 60)
    if hours:
        return f"{hours} h {minutes} min"
    if minutes:
        return f"{minutes} min {secs} s" if secs else f"{minutes} min"
    return f"{secs} s"


def _status_label(value: str | None) -> str:
    return value.replace("_", " ").capitalize() if value else ""


def _summary_source_label(summary: dict) -> str:
    if summary["source"] == "manual":
        return "Written or edited by a user."
    if summary["source"] == "ai":
        return "Generated from the transcript by rule-based extraction."
    return ""


def _detail_rows(meeting: dict) -> list[tuple[str, str]]:
    transcription = _status_label(meeting["transcription_status"])
    if transcription and meeting["transcription_language"]:
        transcription = f"{transcription} ({meeting['transcription_language']})"
    rows = [
        ("Date", _datetime_label(meeting["meeting_date"])),
        ("Duration", _duration_label(meeting["duration_seconds"])),
        ("Platform", meeting["platform"] or ""),
        ("Status", _status_label(meeting["status"])),
        ("Meeting code", meeting["meeting_code"] or ""),
        ("Recording", "Available" if meeting["has_recording"] else "Not available"),
        ("Transcription", transcription),
    ]
    return [(label, value) for label, value in rows if value]


def _segment_heading(segment: dict) -> str:
    start, end = _clock(segment["start_time"]), _clock(segment["end_time"])
    span = f"{start} - {end}" if end and end != start else start
    speaker = _clean(segment["speaker"]).strip() or "Unknown"
    return f"[{span}] {speaker}"


def _generated_label(data: dict) -> str:
    return f"Exported from MeetNote on {_datetime_label(data['generated_at'])}"


# --------------------------------------------------------------------------
# JSON
# --------------------------------------------------------------------------


def _json_default(value):
    if isinstance(value, datetime):
        return value.isoformat().replace("+00:00", "Z")
    if isinstance(value, date):
        return value.isoformat()
    raise TypeError(f"Unsupported value: {type(value).__name__}")


def render_json(data: dict) -> bytes:
    return json.dumps(data, ensure_ascii=False, indent=2, default=_json_default).encode("utf-8")


# --------------------------------------------------------------------------
# Plain text
# --------------------------------------------------------------------------


def _txt_heading(lines: list[str], text: str) -> None:
    lines.extend(["", "", text.upper(), "-" * len(text)])


def _indent(text: str, prefix: str) -> str:
    return text.replace("\n", "\n" + prefix)


def render_txt(data: dict) -> bytes:
    meeting = data["meeting"]
    title = _clean(meeting["title"]).strip() or "Untitled meeting"
    lines = [title, "=" * min(len(title), 80), _generated_label(data)]

    if "details" in data["sections"]:
        _txt_heading(lines, "Meeting details")
        lines.extend(f"{label}: {_clean(value)}" for label, value in _detail_rows(meeting))
        if meeting["description"]:
            lines.extend(["", "Description:", _clean(meeting["description"])])

    if "summary" in data:
        summary = data["summary"]
        _txt_heading(lines, "Summary")
        if summary["text"]:
            lines.append(_clean(summary["text"]))
            lines.extend(["", f"({_summary_source_label(summary)})"] if summary["source"] else [])
        else:
            lines.append("No summary has been saved for this meeting.")
        if summary["key_points"]:
            lines.extend(["", "Key points:"])
            for point in summary["key_points"]:
                stamp = f" [{_clock(point['timestamp'])}]" if point["timestamp"] is not None else ""
                lines.append(f"  - {_indent(_clean(point['text']), '    ')}{stamp}")

    if "action_items" in data:
        _txt_heading(lines, "Action items")
        if not data["action_items"]:
            lines.append("No action items.")
        for number, item in enumerate(data["action_items"], start=1):
            lines.append(f"{number}. [{_status_label(item['status'])}] {_indent(_clean(item['task']), '   ')}")
            extras = [
                f"Assignee: {_clean(item['assignee'])}" if item["assignee"] else "",
                f"Due: {_date_label(item['due_date'])}" if item["due_date"] else "",
                f"Time: {_clock(item['timestamp'])}" if item["timestamp"] is not None else "",
            ]
            extras = [extra for extra in extras if extra]
            if extras:
                lines.append("   " + " | ".join(extras))

    if "decisions" in data:
        _txt_heading(lines, "Decisions")
        if not data["decisions"]:
            lines.append("No decisions.")
        for number, decision in enumerate(data["decisions"], start=1):
            lines.append(f"{number}. [{_clock(decision['timestamp'])}] {_indent(_clean(decision['decision']), '   ')}")
            if decision["context"]:
                lines.append(f"   Context: {_indent(_clean(decision['context']), '   ')}")

    if "transcript" in data:
        _txt_heading(lines, "Transcript")
        if not data["transcript"]:
            lines.append("No transcript.")
        for segment in data["transcript"]:
            lines.append(f"{_segment_heading(segment)}: {_indent(_clean(segment['text']).strip(), '  ')}")

    return ("\n".join(lines).strip() + "\n").encode("utf-8")


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------


def render_docx(data: dict) -> bytes:
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor

    meeting = data["meeting"]
    title = _clean(meeting["title"]).strip() or "Untitled meeting"
    muted = RGBColor(0x64, 0x74, 0x8B)

    document = Document()
    properties = document.core_properties
    properties.title = title
    properties.author = "MeetNote"
    properties.last_modified_by = "MeetNote"
    properties.comments = ""

    document.add_heading(title, level=0)
    _docx_note(document, _generated_label(data), muted, Pt(9))

    def table(headers: list[str] | None, rows: list[list[str]], widths: list[float] | None = None):
        grid = document.add_table(rows=0, cols=len(rows[0]) if rows else len(headers or []))
        grid.style = "Table Grid"
        if headers:
            cells = grid.add_row().cells
            for cell, text in zip(cells, headers):
                cell.paragraphs[0].add_run(text).bold = True
        for values in rows:
            cells = grid.add_row().cells
            for cell, text in zip(cells, values):
                _docx_multiline(cell.paragraphs[0], text or DASH)
        if widths:
            for row in grid.rows:
                for cell, width in zip(row.cells, widths):
                    cell.width = Cm(width)
        return grid

    if "details" in data["sections"]:
        document.add_heading("Meeting details", level=1)
        rows = [[label, _clean(value)] for label, value in _detail_rows(meeting)]
        if rows:
            grid = table(None, rows, [4.0, 12.0])
            for row in grid.rows:
                row.cells[0].paragraphs[0].runs[0].bold = True
        if meeting["description"]:
            document.add_heading("Description", level=2)
            _docx_paragraphs(document, _clean(meeting["description"]))

    if "summary" in data:
        summary = data["summary"]
        document.add_heading("Summary", level=1)
        if summary["text"]:
            _docx_paragraphs(document, _clean(summary["text"]))
            if summary["source"]:
                _docx_note(document, _summary_source_label(summary), muted, Pt(9))
        else:
            _docx_note(document, "No summary has been saved for this meeting.", muted)
        if summary["key_points"]:
            document.add_heading("Key points", level=2)
            for point in summary["key_points"]:
                paragraph = document.add_paragraph(style="List Bullet")
                _docx_multiline(paragraph, _clean(point["text"]))
                if point["timestamp"] is not None:
                    run = paragraph.add_run(f"  [{_clock(point['timestamp'])}]")
                    run.font.color.rgb = muted

    if "action_items" in data:
        document.add_heading("Action items", level=1)
        if data["action_items"]:
            table(
                ["Task", "Assignee", "Due date", "Status", "Time"],
                [
                    [
                        _clean(item["task"]),
                        _clean(item["assignee"]),
                        _date_label(item["due_date"]),
                        _status_label(item["status"]),
                        _clock(item["timestamp"]),
                    ]
                    for item in data["action_items"]
                ],
                [7.5, 3.0, 2.4, 2.1, 1.6],
            )
        else:
            _docx_note(document, "No action items.", muted)

    if "decisions" in data:
        document.add_heading("Decisions", level=1)
        if data["decisions"]:
            table(
                ["Time", "Decision", "Context"],
                [
                    [_clock(decision["timestamp"]), _clean(decision["decision"]), _clean(decision["context"])]
                    for decision in data["decisions"]
                ],
                [1.6, 8.0, 7.0],
            )
        else:
            _docx_note(document, "No decisions.", muted)

    if "transcript" in data:
        document.add_heading("Transcript", level=1)
        if not data["transcript"]:
            _docx_note(document, "No transcript.", muted)
        for segment in data["transcript"]:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(6)
            heading = paragraph.add_run(_segment_heading(segment) + ": ")
            heading.bold = True
            _docx_multiline(paragraph, _clean(segment["text"]).strip())

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _docx_multiline(paragraph, text: str) -> None:
    for index, line in enumerate(text.split("\n")):
        if index:
            paragraph.add_run().add_break()
        paragraph.add_run(line)


def _docx_paragraphs(document, text: str) -> None:
    for block in re.split(r"\n\s*\n", text.strip()):
        _docx_multiline(document.add_paragraph(), block)


def _docx_note(document, text: str, color, size=None) -> None:
    run = document.add_paragraph().add_run(text)
    run.font.color.rgb = color
    if size is not None:
        run.font.size = size


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

_FONT_DIRS = [
    Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts",
    Path("/usr/share/fonts/truetype/dejavu"),
    Path("/usr/share/fonts/dejavu"),
    Path("/usr/share/fonts/TTF"),
    Path("/usr/share/fonts/truetype/noto"),
    Path("/usr/share/fonts/noto"),
    Path("/System/Library/Fonts/Supplemental"),
    Path("/Library/Fonts"),
]
@dataclass(frozen=True)
class _FontFaces:
    regular: Path
    regular_index: int
    bold: Path
    bold_index: int


# Regular and bold faces, in order of preference. A PDF uses the first one
# that can draw every character in the export. A single face per document is
# deliberate: fpdf2 fallback fonts can drop or garble a wrapped line that
# switches fonts, which would lose transcript text.
_FONT_FACES = [
    ("DejaVuSans.ttf", 0, "DejaVuSans-Bold.ttf", 0),
    ("arial.ttf", 0, "arialbd.ttf", 0),
    ("Arial.ttf", 0, "Arial Bold.ttf", 0),
    ("segoeui.ttf", 0, "segoeuib.ttf", 0),
    ("NotoSans-Regular.ttf", 0, "NotoSans-Bold.ttf", 0),
    ("LiberationSans-Regular.ttf", 0, "LiberationSans-Bold.ttf", 0),
    ("Nirmala.ttc", 0, "Nirmala.ttc", 1),
    ("NotoSansDevanagari-Regular.ttf", 0, "NotoSansDevanagari-Bold.ttf", 0),
    ("msyh.ttc", 0, "msyhbd.ttc", 0),
    ("NotoSansCJK-Regular.ttc", 0, "NotoSansCJK-Bold.ttc", 0),
    ("malgun.ttf", 0, "malgunbd.ttf", 0),
    ("Arial Unicode.ttf", 0, "Arial Unicode.ttf", 0),
]
_TYPOGRAPHY = str.maketrans(
    {
        "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
        "\u2013": "-", "\u2014": "-", "\u2022": "-", "\u2026": "...", "\u00a0": " ",
    }
)

# fontTools reports every OpenType table it cannot subset; none affect text.
logging.getLogger("fontTools.subset").setLevel(logging.ERROR)


def _find_font(name: str) -> Path | None:
    candidate = Path(name)
    if candidate.is_absolute():
        return candidate if candidate.is_file() else None
    for directory in _FONT_DIRS:
        candidate = directory / name
        if candidate.is_file():
            return candidate
    return None


def _available_faces() -> list[_FontFaces]:
    specs = list(_FONT_FACES)
    configured = os.environ.get("MEETNOTE_PDF_FONT")
    if configured:
        specs.insert(0, (configured, 0, os.environ.get("MEETNOTE_PDF_FONT_BOLD") or configured, 0))
    faces = []
    for regular, regular_index, bold, bold_index in specs:
        regular_path = _find_font(regular)
        if regular_path is None:
            continue
        bold_path = _find_font(bold)
        if bold_path is None:
            bold_path, bold_index = regular_path, regular_index
        faces.append(_FontFaces(regular_path, regular_index, bold_path, bold_index))
    return faces


@lru_cache(maxsize=32)
def _font_charset(path: Path, index: int) -> frozenset[int]:
    from fontTools.ttLib import TTFont

    try:
        with TTFont(str(path), fontNumber=index, lazy=True) as font:
            return frozenset(font.getBestCmap() or {})
    except Exception:
        logger.warning("Could not read font %s for PDF export", path.name)
        return frozenset()


def _choose_font(characters: set[int]) -> tuple[_FontFaces, frozenset[int]] | None:
    needed = {code for code in characters if code > 0x20 and not chr(code).isspace()}
    best = None
    best_missing = 0
    for face in _available_faces():
        charset = _font_charset(face.regular, face.regular_index)
        if not charset:
            continue
        missing = len(needed - charset)
        if missing == 0:
            return face, charset
        if best is None or missing < best_missing:
            best, best_missing = (face, charset), missing
    return best


def _document_characters(data: dict) -> set[int]:
    characters: set[int] = set()

    def visit(value) -> None:
        if isinstance(value, str):
            characters.update(map(ord, _clean(value)))
        elif isinstance(value, dict):
            for item in value.values():
                visit(item)
        elif isinstance(value, list):
            for item in value:
                visit(item)

    visit(data)
    return characters


def _make_pdf(data: dict):
    from fpdf import FPDF

    class ReportPDF(FPDF):
        body_family = "Helvetica"
        charset: frozenset[int] | None = None

        def footer(self) -> None:
            self.set_y(-12)
            self.set_font(self.body_family, "", 8)
            self.set_text_color(120, 120, 120)
            self.cell(0, 5, f"Page {self.page_no()} of {{nb}}", align="C")

        def text_for(self, text: str) -> str:
            text = _clean(text)
            if self.charset is not None:
                return text
            return text.translate(_TYPOGRAPHY).encode("latin-1", "replace").decode("latin-1")

        def symbol(self, preferred: str, plain: str) -> str:
            return preferred if self.charset is not None and ord(preferred) in self.charset else plain

    def blank() -> ReportPDF:
        pdf = ReportPDF(format="A4", unit="mm")
        pdf.set_margins(18, 18, 18)
        pdf.set_auto_page_break(auto=True, margin=18)
        return pdf

    pdf = blank()
    chosen = _choose_font(_document_characters(data))
    if chosen is None:
        logger.warning("No Unicode font found for PDF export; using Helvetica")
        return pdf
    face, charset = chosen
    try:
        pdf.add_font("Body", "", str(face.regular), collection_font_number=face.regular_index)
        pdf.add_font("Body", "B", str(face.bold), collection_font_number=face.bold_index)
    except Exception:
        logger.warning("Could not load font %s for PDF export; using Helvetica", face.regular.name)
        return blank()
    pdf.body_family = "Body"
    pdf.charset = charset
    return pdf


# Longest cell text drawn as a table row. A row cannot span pages, so longer
# entries are drawn as flowing paragraphs instead.
_MAX_TABLE_CELL = 1200


def render_pdf(data: dict) -> bytes:
    from fpdf.fonts import FontFace

    pdf = _make_pdf(data)
    family = pdf.body_family
    t = pdf.text_for
    dash = pdf.symbol(DASH, "-")
    bullet = pdf.symbol("\u2022", "-")
    width = pdf.epw
    ink = (15, 23, 42)
    muted = (100, 116, 139)
    rule = (203, 213, 225)

    meeting = data["meeting"]
    title = meeting["title"].strip() if meeting["title"] else "Untitled meeting"
    pdf.set_title(_clean(title))
    pdf.set_author("MeetNote")
    pdf.set_creator("MeetNote")
    pdf.add_page()

    def font(size: float, bold: bool = False, color=ink) -> None:
        pdf.set_font(family, "B" if bold else "", size)
        pdf.set_text_color(*color)

    def paragraph(text: str, size: float = 10.5, color=ink, bold: bool = False, height: float = 5.4) -> None:
        font(size, bold, color)
        pdf.multi_cell(width, height, t(text), new_x="LMARGIN", new_y="NEXT")

    def heading(text: str, level: int = 1) -> None:
        size = 14 if level == 1 else 11.5
        if pdf.will_page_break(size + 14):
            pdf.add_page()
        pdf.ln(5 if level == 1 else 3)
        font(size, True)
        pdf.cell(width, size * 0.55, t(text), new_x="LMARGIN", new_y="NEXT")
        if level == 1:
            pdf.set_draw_color(*rule)
            pdf.set_line_width(0.3)
            pdf.line(pdf.l_margin, pdf.get_y() + 1, pdf.l_margin + width, pdf.get_y() + 1)
            pdf.ln(3)
        else:
            pdf.ln(1)

    def table(
        headers: list[str] | None,
        rows: list[list[str]],
        col_widths: tuple[float, ...],
        bold_first_column: bool = False,
    ) -> None:
        label_style = FontFace(emphasis="BOLD", color=ink)
        font(9.5)
        pdf.set_draw_color(*rule)
        with pdf.table(
            col_widths=col_widths,
            width=width,
            text_align="LEFT",
            line_height=5,
            padding=1.6,
            first_row_as_headings=bool(headers),
            repeat_headings=1,
            headings_style=FontFace(emphasis="BOLD", color=ink, fill_color=(241, 245, 249)),
            borders_layout="HORIZONTAL_LINES",
        ) as grid:
            if headers:
                header_row = grid.row()
                for value in headers:
                    header_row.cell(t(value))
            for values in rows:
                row = grid.row()
                for column, value in enumerate(values):
                    style = label_style if bold_first_column and column == 0 else None
                    row.cell(t(value) if value else dash, style=style)
        pdf.ln(2)

    def fits_table(rows: list[list[str]]) -> bool:
        return all(len(value or "") <= _MAX_TABLE_CELL for values in rows for value in values)

    font(20, True)
    pdf.multi_cell(width, 9, t(title), new_x="LMARGIN", new_y="NEXT")
    paragraph(_generated_label(data), size=9, color=muted)

    if "details" in data["sections"]:
        heading("Meeting details")
        rows = [[label, value] for label, value in _detail_rows(meeting)]
        if rows:
            table(None, rows, (1, 3), bold_first_column=True)
        if meeting["description"]:
            heading("Description", level=2)
            paragraph(meeting["description"])

    if "summary" in data:
        summary = data["summary"]
        heading("Summary")
        if summary["text"]:
            for block in re.split(r"\n\s*\n", _clean(summary["text"]).strip()):
                paragraph(block)
                pdf.ln(1.5)
            if summary["source"]:
                paragraph(_summary_source_label(summary), size=8.5, color=muted)
        else:
            paragraph("No summary has been saved for this meeting.", color=muted)
        if summary["key_points"]:
            heading("Key points", level=2)
            for point in summary["key_points"]:
                stamp = f"  [{_clock(point['timestamp'])}]" if point["timestamp"] is not None else ""
                font(10.5)
                pdf.set_x(pdf.l_margin)
                pdf.cell(5, 5.4, bullet)
                pdf.multi_cell(width - 5, 5.4, t(point["text"] + stamp), new_x="LMARGIN", new_y="NEXT")
                pdf.ln(0.8)

    if "action_items" in data:
        heading("Action items")
        items = data["action_items"]
        rows = [
            [
                _clean(item["task"]),
                _clean(item["assignee"]),
                _date_label(item["due_date"]),
                _status_label(item["status"]),
                _clock(item["timestamp"]),
            ]
            for item in items
        ]
        if not items:
            paragraph("No action items.", color=muted)
        elif fits_table(rows):
            table(["Task", "Assignee", "Due date", "Status", "Time"], rows, (46, 18, 15, 13, 9))
        else:
            for number, (task, assignee, due, status_value, time) in enumerate(rows, start=1):
                paragraph(f"{number}. {task}", bold=True)
                meta = " | ".join(
                    value
                    for value in (
                        f"Status: {status_value}",
                        f"Assignee: {assignee}" if assignee else "",
                        f"Due: {due}" if due else "",
                        f"Time: {time}" if time else "",
                    )
                    if value
                )
                paragraph(meta, size=9, color=muted)
                pdf.ln(2)

    if "decisions" in data:
        heading("Decisions")
        decisions = data["decisions"]
        rows = [
            [_clock(decision["timestamp"]), _clean(decision["decision"]), _clean(decision["context"])]
            for decision in decisions
        ]
        if not decisions:
            paragraph("No decisions.", color=muted)
        elif fits_table(rows):
            table(["Time", "Decision", "Context"], rows, (9, 50, 41))
        else:
            for number, (time, decision_text, context) in enumerate(rows, start=1):
                paragraph(f"{number}. [{time}] {decision_text}", bold=True)
                if context:
                    paragraph(f"Context: {context}", size=9, color=muted)
                pdf.ln(2)

    if "transcript" in data:
        heading("Transcript")
        if not data["transcript"]:
            paragraph("No transcript.", color=muted)
        for segment in data["transcript"]:
            # Keep a speaker line together with the start of its text.
            if pdf.will_page_break(12):
                pdf.add_page()
            paragraph(_segment_heading(segment), size=9, bold=True, color=(51, 65, 85), height=4.8)
            paragraph(_clean(segment["text"]).strip(), size=10.5, height=5.4)
            pdf.ln(2.2)

    return bytes(pdf.output())


_RENDERERS = {
    ExportFormat.PDF: render_pdf,
    ExportFormat.DOCX: render_docx,
    ExportFormat.TXT: render_txt,
    ExportFormat.JSON: render_json,
}
